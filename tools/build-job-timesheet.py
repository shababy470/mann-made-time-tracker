#!/usr/bin/env python3
"""Build a client-facing timesheet for one job from the MANNMADE Time Log.

Written for MM02755 (Liberty — Challenge 2026) when the client asked for a
timesheet, and kept because the next client will ask too.

    python3 tools/build-job-timesheet.py <time-log.xlsx> <JOB NUMBER> [out.xlsx]

`time-log.xlsx` is the "MANNMADE Time Log" spreadsheet exported from Drive as
xlsx (id 1yZoZ_mfeEP37gZ1cfVFOVOJqB3a3hw1JyhtfIvJRV2U).

Two things it does that matter, and why:

  * Combines people who logged under two spellings of their own name. The Time
    Log matches on typed name, so "Josh Lindberg" and "josh lindberg" are two
    people to the sheet and one person to the client. See MERGE below — extend
    it when a new pair appears.
  * Drops exact duplicate entries (same person, date, start time and duration).
    These come from the calendar import, whose de-duplication is broken: it
    compares a Date object against an "HH:MM" string, so it has never matched.
    Billing a client for hours logged twice is not a rounding error, so they
    come out — and land on a "Removed" tab so the decision is visible.

Totals are written as values, not formulas, and cross-checked three ways before
the file is saved. LibreOffice cannot recalculate in the sandbox this runs in,
so a formula-based workbook could not be verified and would open blank in
anything reading cached values.
"""
import sys, datetime, collections
import openpyxl
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side

# Time Log names that belong to the same person. Left spelling -> right spelling.
MERGE = {
    'josh lindberg': 'Josh Lindberg',
    'Tshepo': 'Tshepo Kgame',
    'Brad Longbottom': 'Bradley Longbottom',
}

ARIAL = 'Arial'
NAVY = PatternFill('solid', fgColor='0A0A0A')
HDR = Font(name=ARIAL, bold=True, color='FFFFFF', size=10)
_thin = Side(style='thin', color='D9D9D9')
BOX = Border(left=_thin, right=_thin, top=_thin, bottom=_thin)


def load_job(path, job):
    """Every Time Log row for one job, name-merged, split into kept and duplicates.

    Returns (kept, removed, job_name) — job_name is the most common Job Name
    recorded against the number, since people type it slightly differently.
    """
    ws = openpyxl.load_workbook(path, data_only=True)['Time Log']
    recs = []
    names = collections.Counter()
    for r in list(ws.iter_rows(values_only=True))[1:]:
        if str(r[2] or '').strip().upper() != job.upper():
            continue
        d = r[0].date() if isinstance(r[0], datetime.datetime) else r[0]
        t = r[6]
        start = t.strftime('%H:%M') if hasattr(t, 'strftime') else str(t)[:5]
        who = str(r[1] or '').strip()
        if r[3]:
            names[str(r[3]).strip()] += 1
        recs.append({
            'date': d, 'person': MERGE.get(who, who), 'task': str(r[5] or '').strip(),
            'start': start, 'hrs': round(float(r[8] or 0), 2),
            'bill': 'Yes' if str(r[9]).strip() == 'Yes' else 'No',
        })
    seen, kept, removed = set(), [], []
    for x in sorted(recs, key=lambda z: (z['date'], z['person'], z['start'])):
        k = (x['person'], x['date'], x['start'], x['hrs'])
        (kept if k not in seen else removed).append(x)
        seen.add(k)
    job_name = names.most_common(1)[0][0] if names else ''
    return kept, removed, job_name


def totals(kept):
    tot = round(sum(x['hrs'] for x in kept), 2)
    bil = round(sum(x['hrs'] for x in kept if x['bill'] == 'Yes'), 2)
    non = round(sum(x['hrs'] for x in kept if x['bill'] == 'No'), 2)
    byp = collections.Counter()
    bym = collections.Counter()
    for x in kept:
        byp[x['person']] += x['hrs']
        bym[(x['date'].year, x['date'].month)] += x['hrs']
    byp = {k: round(v, 2) for k, v in byp.items()}
    bym = {k: round(v, 2) for k, v in bym.items()}
    # The whole document is these numbers, so prove they reconcile before saving.
    assert abs(bil + non - tot) < 0.02, (bil, non, tot)
    assert abs(sum(byp.values()) - tot) < 0.05, (sum(byp.values()), tot)
    assert abs(sum(bym.values()) - tot) < 0.05, (sum(bym.values()), tot)
    return tot, bil, non, byp, bym


def write(path, job, title, kept, removed):
    tot, bil, non, byp, bym = totals(kept)
    wb = openpyxl.Workbook()

    det = wb.active
    det.title = 'Detail'
    for i, c in enumerate(['Date', 'Team Member', 'Description of Work', 'Hours', 'Billable'], 1):
        cell = det.cell(row=1, column=i, value=c)
        cell.font, cell.fill = HDR, NAVY
    det.freeze_panes = 'A2'
    for i, x in enumerate(kept, start=2):
        det.cell(row=i, column=1, value=x['date']).number_format = 'yyyy-mm-dd'
        det.cell(row=i, column=2, value=x['person'])
        det.cell(row=i, column=3, value=x['task'] or '—')
        det.cell(row=i, column=4, value=x['hrs']).number_format = '0.00'
        det.cell(row=i, column=5, value=x['bill'])
        for col in range(1, 6):
            c = det.cell(row=i, column=col)
            c.font, c.border = Font(name=ARIAL, size=10), BOX
    last = len(kept) + 1
    det.cell(row=last + 1, column=3, value='TOTAL').font = Font(name=ARIAL, bold=True, size=10)
    c = det.cell(row=last + 1, column=4, value=tot)
    c.font, c.number_format = Font(name=ARIAL, bold=True, size=10), '0.00'
    for w, col in zip([12, 24, 64, 10, 10], 'ABCDE'):
        det.column_dimensions[col].width = w
    det.auto_filter.ref = 'A1:E%d' % last

    s = wb.create_sheet('Summary', 0)
    s['A1'] = title
    s['A1'].font = Font(name=ARIAL, bold=True, size=16)
    s['A2'] = 'Timesheet prepared for the client'
    s['A2'].font = Font(name=ARIAL, size=11, color='595959')
    r = 4
    meta = [('Agency', 'MANNMADE'), ('Job number', job),
            ('Period covered', '%s to %s' % (min(x['date'] for x in kept).strftime('%d %b %Y'),
                                             max(x['date'] for x in kept).strftime('%d %b %Y'))),
            ('Prepared', datetime.date.today().strftime('%d %b %Y')),
            ('Entries', len(kept))]
    for k, v in meta:
        s.cell(row=r, column=1, value=k).font = Font(name=ARIAL, bold=True, size=10)
        s.cell(row=r, column=2, value=v).font = Font(name=ARIAL, size=10)
        r += 1
    r += 1
    s.cell(row=r, column=1, value='TOTAL HOURS').font = Font(name=ARIAL, bold=True, size=11)
    c = s.cell(row=r, column=2, value=tot)
    c.font, c.number_format = Font(name=ARIAL, bold=True, size=14), '0.00'
    r += 1
    for lbl, val in [('Billable hours', bil), ('Non-billable hours', non)]:
        s.cell(row=r, column=1, value=lbl).font = Font(name=ARIAL, size=10)
        c = s.cell(row=r, column=2, value=val)
        c.font, c.number_format = Font(name=ARIAL, size=10), '0.00'
        r += 1

    r += 1
    s.cell(row=r, column=1, value='HOURS BY TEAM MEMBER').font = Font(name=ARIAL, bold=True, size=11)
    r += 1
    for i, h in enumerate(['Team Member', 'Hours', '% of job'], 1):
        c = s.cell(row=r, column=i, value=h)
        c.font, c.fill = HDR, NAVY
    r += 1
    for n, v in sorted(byp.items(), key=lambda kv: -kv[1]):
        s.cell(row=r, column=1, value=n)
        s.cell(row=r, column=2, value=v).number_format = '0.00'
        s.cell(row=r, column=3, value=round(v / tot, 4)).number_format = '0.0%'
        for col in (1, 2, 3):
            c = s.cell(row=r, column=col)
            c.font, c.border = Font(name=ARIAL, size=10), BOX
        r += 1
    s.cell(row=r, column=1, value='Total').font = Font(name=ARIAL, bold=True, size=10)
    c = s.cell(row=r, column=2, value=round(sum(byp.values()), 2))
    c.font, c.number_format = Font(name=ARIAL, bold=True, size=10), '0.00'
    r += 2

    s.cell(row=r, column=1, value='HOURS BY MONTH').font = Font(name=ARIAL, bold=True, size=11)
    r += 1
    for i, h in enumerate(['Month', 'Hours'], 1):
        c = s.cell(row=r, column=i, value=h)
        c.font, c.fill = HDR, NAVY
    r += 1
    for (y, m), v in sorted(bym.items()):
        s.cell(row=r, column=1, value=datetime.date(y, m, 1).strftime('%b %Y'))
        s.cell(row=r, column=2, value=v).number_format = '0.00'
        for col in (1, 2):
            c = s.cell(row=r, column=col)
            c.font, c.border = Font(name=ARIAL, size=10), BOX
        r += 1
    s.cell(row=r, column=1, value='Total').font = Font(name=ARIAL, bold=True, size=10)
    c = s.cell(row=r, column=2, value=round(sum(bym.values()), 2))
    c.font, c.number_format = Font(name=ARIAL, bold=True, size=10), '0.00'
    r += 2

    s.cell(row=r, column=1, value='NOTES ON THIS TIMESHEET').font = Font(name=ARIAL, bold=True, size=11)
    r += 1
    notes = [
        "Source: MANNMADE's Time Log — entries recorded by each team member as the work was done, "
        "against job %s." % job,
        'Figures on this tab are a snapshot of the Detail tab as at the date above. If rows on '
        'Detail are changed, this summary does not update automatically.',
        '%d duplicate entries totalling %.2f hours were identified and excluded — the same person, '
        'date, start time and duration recorded twice. They are listed on the "Removed" tab for '
        'transparency.' % (len(removed), sum(x['hrs'] for x in removed)),
        'Where a team member had logged time under two spellings of their name, those entries have '
        'been combined.',
        'Non-billable hours are shown separately above and flagged in the Billable column on the '
        'Detail tab.',
    ]
    for n in notes:
        c = s.cell(row=r, column=1, value='• ' + n)
        c.font = Font(name=ARIAL, size=9, color='595959')
        c.alignment = Alignment(wrap_text=True, vertical='top')
        s.merge_cells(start_row=r, start_column=1, end_row=r, end_column=3)
        s.row_dimensions[r].height = 30
        r += 1
    for w, col in zip([36, 14, 12], 'ABC'):
        s.column_dimensions[col].width = w

    rm = wb.create_sheet('Removed')
    for i, h in enumerate(['Date', 'Team Member', 'Description of Work', 'Hours', 'Why excluded'], 1):
        c = rm.cell(row=1, column=i, value=h)
        c.font, c.fill = HDR, NAVY
    for i, x in enumerate(removed, start=2):
        rm.cell(row=i, column=1, value=x['date']).number_format = 'yyyy-mm-dd'
        rm.cell(row=i, column=2, value=x['person'])
        rm.cell(row=i, column=3, value=x['task'] or '—')
        rm.cell(row=i, column=4, value=x['hrs']).number_format = '0.00'
        rm.cell(row=i, column=5, value='Duplicate of an identical entry on the same date and start time')
        for col in range(1, 6):
            rm.cell(row=i, column=col).font = Font(name=ARIAL, size=10)
    for w, col in zip([12, 24, 52, 10, 54], 'ABCDE'):
        rm.column_dimensions[col].width = w

    wb.save(path)
    return tot, bil, non, len(kept), len(removed)


def main():
    if len(sys.argv) < 3:
        sys.exit(__doc__)
    src, job = sys.argv[1], sys.argv[2]
    out = sys.argv[3] if len(sys.argv) > 3 else '%s - Timesheet.xlsx' % job
    kept, removed, job_name = load_job(src, job)
    if not kept:
        sys.exit('No Time Log entries found for %s.' % job)
    title = '%s · %s' % (job, job_name) if job_name else job
    tot, bil, non, nk, nr = write(out, job, title, kept, removed)
    print('%s: %d entries, %.2f h (%.2f billable, %.2f non-billable); '
          '%d duplicate rows excluded. Saved to %s' % (job, nk, tot, bil, non, nr, out))


if __name__ == '__main__':
    main()
